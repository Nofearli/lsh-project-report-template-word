# -*- coding: utf-8 -*-
"""按《项目报告模板》格式生成 Word 项目报告。

用法（由模型按报告内容编写驱动脚本，import 本模块）：

    import sys
    sys.path.insert(0, r"<skill目录>/scripts")
    from build_report import Report, cn_num

    r = Report()
    r.set_cover(title="XX系统测试报告", date="2026年 10月")
    r.h1(f"{cn_num(1)}.概述")
    r.para("本章介绍……")
    r.h2("1.1 背景")
    r.para("正文段落，自动首行缩进……")
    r.figure(r"img\fig1.png", "系统架构图", width_cm=12)
    r.table(caption="测试结果", headers=["项目", "结果"], rows=[["响应时间", "1.2 s"]])
    r.references(["作者. 题名[J]. 期刊, 2024, 40(2): 15-20."])
    r.save(r"输出.docx")

模板底版 assets/template-base.docx 已内置全部样式（正文/标题/题注/页码分节等），
本模块只负责向其中填充内容，保证与模板格式一致。格式细节见 references/format-spec.md。
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor, Twips

SKILL_DIR = Path(__file__).resolve().parents[1]
TEMPLATE = SKILL_DIR / "assets" / "template-base.docx"

# 页面可用宽度（A4 21cm - 左右 3.18cm×2 ≈ 14.65cm），图片宽度上限
MAX_IMG_WIDTH_CM = 14.5

_CN = "零一二三四五六七八九"


def cn_num(n: int) -> str:
    """1->一, 10->十, 12->十二, 20->二十, 21->二十一。用于一级标题编号。"""
    if n <= 0:
        return str(n)
    if n < 10:
        return _CN[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + _CN[n % 10]
    tens, ones = divmod(n, 10)
    return _CN[tens] + "十" + (_CN[ones] if ones else "")


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_pstyle(p, style_id: str):
    """按 styleId 设置段落样式（模板样式 id：1/2/3=标题，af0=题注，af2=图表目录项…）。"""
    pPr = p._p.get_or_add_pPr()
    el = pPr.find(qn("w:pStyle"))
    if el is None:
        el = _el("w:pStyle")
        pPr.insert(0, el)
    el.set(qn("w:val"), style_id)


def set_first_line_chars(p, chars=200, twips=480):
    """首行缩进按"字符"设置（随字号自适应），并给出 twips 兜底值。"""
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = _el("w:ind")
        jc = pPr.find(qn("w:jc"))
        if jc is not None:
            jc.addprevious(ind)
        else:
            pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), str(twips))


def add_field_char(p, fld_type: str):
    r = p.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": fld_type}))
    return r


def add_instr_text(p, instr: str):
    r = p.add_run()
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._r.append(it)
    return r


def add_seq_field(p, identifier: str, cached: str = "1"):
    """插入 SEQ 编号域（题注自动编号）。"""
    add_field_char(p, "begin")
    add_instr_text(p, f" SEQ {identifier} \\* ARABIC ")
    add_field_char(p, "separate")
    p.add_run(cached)
    add_field_char(p, "end")


class Report:
    """面向内容的报告生成器。所有内容追加到正文节（第 3 节）末尾。"""

    def __init__(self, template: str | Path = TEMPLATE):
        self.doc = Document(str(template))
        self._seq_counters = {"图": 0, "表": 0}
        self._insert_toc_fields()
        self._ensure_update_fields()

    # ------------------------------------------------------------------ 封面
    def set_cover(self, title: str, org: str = "北京航空航天大学",
                  date: str | None = None):
        """填写封面标题、落款单位与日期。标题内可用 \\n 分两行。"""
        if date is None:
            now = _dt.date.today()
            date = f"{now.year}年 {now.month}月"
        tbl = self.doc.tables[0]
        title_cell = tbl.rows[1].cells[0]
        lines = title.split("\n")
        paras = title_cell.paragraphs  # 模板：2 个标题段 + 1 个空段
        for i, p in enumerate(paras[:2]):
            text = lines[i] if i < len(lines) else ""
            self._replace_para_text(p, text, sz=52)
        date_cell = tbl.rows[8].cells[0]
        dparas = date_cell.paragraphs  # 模板：2 空段 + 单位段 + 日期段
        self._replace_para_text(dparas[2], org, sz=32)
        self._replace_para_text(dparas[3], date, sz=32)

    @staticmethod
    def _replace_para_text(p, text: str, sz: int):
        """清空段落 run 后按模板格式重写（黑体、指定半磅字号）。"""
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        if not text:
            return
        run = p.add_run(text)
        rPr = run._r.get_or_add_rPr()
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = _el("w:rFonts")
            rPr.insert(0, fonts)
        fonts.set(qn("w:eastAsia"), "黑体")
        fonts.set(qn("w:hint"), "eastAsia")
        rPr.append(_el("w:kern", **{"w:val": 0}))
        rPr.append(_el("w:sz", **{"w:val": sz}))
        rPr.append(_el("w:szCs", **{"w:val": sz}))

    # ------------------------------------------------------------------ 目录
    def _insert_toc_fields(self):
        """把底版中的占位段替换为 TOC 域（打开文档后自动/手动刷新生成条目）。"""
        fields = {
            "[[TOC_MAIN]]": ' TOC \\o "1-3" \\h \\z \\u ',
            "[[TOC_FIG]]": ' TOC \\h \\z \\c "图" ',
            "[[TOC_TAB]]": ' TOC \\h \\z \\c "表" ',
        }
        for p in self.doc.paragraphs:
            text = p.text
            for marker, instr in fields.items():
                if marker in text:
                    for r in list(p.runs):
                        r._r.getparent().remove(r._r)
                    add_field_char(p, "begin")
                    add_instr_text(p, instr)
                    add_field_char(p, "separate")
                    hint = p.add_run("（此目录为域，打开文档后自动更新；"
                                     "若未刷新请全选按 F9 或右键“更新域”）")
                    hint.font.color.rgb = RGBColor.from_string("808080")
                    hint.font.size = Pt(9)
                    hint.font.italic = True
                    add_field_char(p, "end")

    def _ensure_update_fields(self):
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            uf = parse_xml('<w:updateFields %s w:val="true"/>' % nsdecls("w"))
            compat = settings.find(qn("w:compat"))
            if compat is not None:
                compat.addprevious(uf)
            else:
                settings.append(uf)

    def remove_list_of_figures(self):
        """报告无图时，删除"图片"图目录块（标题段 + 域段）。"""
        self._remove_front_block("图片", "[[TOC_FIG]]")

    def remove_list_of_tables(self):
        """报告无表时，删除"表格"表目录块（标题段 + 域段）。"""
        self._remove_front_block("表格", "[[TOC_TAB]]")

    def _remove_front_block(self, title, marker):
        doomed = []
        seen_title = False
        for p in self.doc.paragraphs:
            if p.text.strip() == title:
                seen_title = True
                doomed.append(p)
            elif marker in p.text and seen_title:
                doomed.append(p)
                break
        for p in doomed:
            p._p.getparent().remove(p._p)

    # ------------------------------------------------------------------ 标题
    def h1(self, text: str):
        """一级标题：黑体+Times New Roman 16pt 加粗居中（编号如“一.”写入 text）。"""
        p = self.doc.add_paragraph()
        set_pstyle(p, "1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(text)
        return p

    def h2(self, text: str):
        """二级标题：宋体+Times New Roman 16pt 加粗（编号如“1.1”写入 text）。"""
        p = self.doc.add_paragraph()
        set_pstyle(p, "2")
        p.add_run(text)
        return p

    def h3(self, text: str):
        """三级标题：宋体+Times New Roman 14pt 加粗（编号如“1.1.1”写入 text）。"""
        p = self.doc.add_paragraph()
        set_pstyle(p, "3")
        p.add_run(text)
        return p

    # ------------------------------------------------------------------ 正文
    def para(self, text: str = "", indent: bool = True,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """正文段落：12pt Times New Roman/宋体、1.5 倍行距、两端对齐。"""
        p = self.doc.add_paragraph()
        if indent:
            set_first_line_chars(p)
        if text:
            p.add_run(text)
        if align is not None:
            p.alignment = align
        return p

    def para_rich(self, segments, indent: bool = True):
        """带格式的正文段落。segments: str 或 (str, fmt) 列表。

        fmt 可用键：bold/italic/sup/sub（布尔）、east（中文字体）、
        ascii（西文字体）、size（磅）、color（如 "808080"）。
        上标引用示例：para_rich(["测试结果满足要求", ("[1]", {"sup": True})])
        """
        p = self.doc.add_paragraph()
        if indent:
            set_first_line_chars(p)
        for seg in segments:
            text, fmt = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run(text)
            if fmt.get("bold"):
                r.bold = True
            if fmt.get("italic"):
                r.italic = True
            if fmt.get("sup"):
                r.font.superscript = True
            if fmt.get("sub"):
                r.font.subscript = True
            if fmt.get("size"):
                r.font.size = Pt(fmt["size"])
            if fmt.get("color"):
                r.font.color.rgb = _rgb(fmt["color"])
            fonts = r._r.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is None and (fmt.get("east") or fmt.get("ascii")):
                fonts = _el("w:rFonts")
                r._r.get_or_add_rPr().insert(0, fonts)
            if fmt.get("east"):
                fonts.set(qn("w:eastAsia"), fmt["east"])
            if fmt.get("ascii"):
                fonts.set(qn("w:ascii"), fmt["ascii"])
                fonts.set(qn("w:hAnsi"), fmt["ascii"])
        return p

    def page_break(self):
        """手动分页（模板在目录与图表目录之间即如此分页）。"""
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.append(_el("w:widowControl"))
        pPr.append(_el("w:spacing", **{"w:line": 240, "w:lineRule": "auto"}))
        pPr.append(_el("w:jc", **{"w:val": "left"}))
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    # ------------------------------------------------------------------ 题注
    def _caption(self, kind: str, text: str, keep_next: bool):
        p = self.doc.add_paragraph()
        set_pstyle(p, "af0")  # 题注：宋体 10.5pt 居中
        self._seq_counters[kind] += 1
        p.add_run(f"{kind} ")
        # 缓存值写真实序号：只做一次域更新（如打开时自动更新）时，
        # 图表目录条目就能取到正确编号
        add_seq_field(p, kind, cached=str(self._seq_counters[kind]))
        p.add_run(f" {text}")
        if keep_next:
            p.paragraph_format.keep_with_next = True
        return p

    # ------------------------------------------------------------------ 图
    def figure(self, image_path: str, caption: str, width_cm: float | None = None):
        """居中插图 + 下方“图 N”题注（SEQ 域自动编号）。宽度默认不超过版心。"""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(str(image_path), width=Cm(width_cm) if width_cm else None)
        shape = self.doc.inline_shapes[-1]
        if width_cm is None and shape.width > Cm(MAX_IMG_WIDTH_CM):
            scale = Cm(MAX_IMG_WIDTH_CM) / shape.width
            shape.width = int(shape.width * scale)
            shape.height = int(shape.height * scale)
        return self._caption("图", caption, keep_next=False)

    # ------------------------------------------------------------------ 表
    def table(self, caption: str, headers=None, rows=None, kind: str = "grid",
              widths_cm=None, header_shading: bool = True,
              header_bold: bool = True, cell_align="center"):
        """插入表格 + 上方“表 N”题注。

        kind="grid"      全框线 0.5pt（Table Grid，数据表，模板表 5-9 样式）
        kind="threeline" 上下 1.5pt 粗线 + 内部 0.5pt 细线、无左右外框
                         （框架/流程表，模板表 4 样式；自动加粗表头）
        widths_cm        各列宽（cm）；省略则自动
        """
        if caption:
            self._caption("表", caption, keep_next=True)
        data = ([headers] if headers else []) + list(rows or [])
        if not data:
            raise ValueError("表格至少要有一行内容")
        ncols = max(len(r) for r in data)
        t = self.doc.add_table(rows=0, cols=ncols)
        if kind == "grid":
            t.style = self.doc.styles["Table Grid"]  # styleId aa
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
        elif kind == "threeline":
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            self._threeline_borders(t)
        else:
            raise ValueError(f"未知表格类型：{kind}")
        if widths_cm:
            t.autofit = False
            for i, w in enumerate(widths_cm):
                t.columns[i].width = Cm(w)

        for ri, row_data in enumerate(data):
            row = t.add_row()
            if kind == "grid":
                trPr = row._tr.get_or_add_trPr()
                trPr.append(_el("w:trHeight", **{"w:val": 433,
                                                 "w:hRule": "atLeast"}))
            if ri == 0:
                row._tr.get_or_add_trPr().append(_el("w:tblHeader"))
            row._tr.get_or_add_trPr().append(_el("w:cantSplit"))
            is_header = ri == 0 and headers is not None
            for ci in range(ncols):
                cell = row.cells[ci]
                text = str(row_data[ci]) if ci < len(row_data) else ""
                # 表头居中；数据格默认居中，cell_align="left" 时左对齐
                if is_header or cell_align == "center":
                    align = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    align = WD_ALIGN_PARAGRAPH.LEFT
                self._fill_cell(cell, text, align,
                                bold=is_header and header_bold)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if is_header and header_shading and kind == "grid":
                    self._shade(cell, "F3F3F3")
                if is_header and kind == "threeline":
                    self._header_bottom_border(cell)
        return t

    @staticmethod
    def _fill_cell(cell, text, align, bold=False):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = align
        run = p.add_run(text)
        if bold:
            run.bold = True

    @staticmethod
    def _shade(cell, fill: str):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto",
                                    "w:fill": fill}))

    @staticmethod
    def _threeline_borders(t):
        tblPr = t._tbl.tblPr
        borders = _el("w:tblBorders")
        for edge, sz in (("top", 12), ("bottom", 12),
                         ("insideH", 4), ("insideV", 4)):
            borders.append(_el(f"w:{edge}", **{"w:val": "single",
                                               "w:sz": sz, "w:space": 0,
                                               "w:color": "auto"}))
        tblPr.append(borders)

    @staticmethod
    def _header_bottom_border(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = _el("w:tcBorders")
        tcBorders.append(_el("w:bottom", **{"w:val": "single", "w:sz": 12,
                                            "w:space": 0, "w:color": "auto"}))
        tcPr.append(tcBorders)

    # ------------------------------------------------------------------ 参考文献
    def references(self, entries, heading: str = "参考文献"):
        """GB/T 7714 顺序编码制文末列表。正文中的上标引用用
        para_rich([..., ("[1]", {"sup": True})]) 自行标注。"""
        self.h1(heading)
        for i, entry in enumerate(entries, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pPr = p._p.get_or_add_pPr()
            # 悬挂缩进 2 字符：编号 [N] 顶格，续行与正文对齐
            ind = _el("w:ind", **{"w:leftChars": 200, "w:left": 480,
                                  "w:hangingChars": 200, "w:hanging": 480})
            pPr.append(ind)
            p.add_run(f"[{i}] {entry}")
        return self

    # ------------------------------------------------------------------ 保存
    def save(self, path: str | Path):
        self.doc.save(str(path))
        return Path(path)


if __name__ == "__main__":
    # 最小示例：生成一份示例报告到当前目录
    r = Report()
    r.set_cover(title="示例项目报告")
    r.h1("一.概述")
    r.para("这是一份按模板格式生成的示例段落。")
    r.h2("1.1 测试数据")
    r.table(caption="示例数据表", headers=["项目", "结果"],
            rows=[["响应时间", "1.2 s"], ["成功率", "99.9%"]])
    r.h1("二.结论")
    r.para_rich(["测试全部通过", ("[1]", {"sup": True}), "。"])
    r.references(["张三. 某系统设计[J]. 计算机工程, 2024, 50(2): 1-10."])
    out = r.save("示例报告.docx")
    print("written:", out)
